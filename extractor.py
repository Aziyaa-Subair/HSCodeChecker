"""
extractor.py
------------
Finds HS-code candidates inside documents so the user doesn't have to
retype them.

Pipeline:
    file (PDF / Excel / Word / CSV / TXT / image)
        -> pull out text (and table cells)
        -> find numbers that look like HS codes
        -> score how confident we are in each one
        -> the GUI shows them on a review screen, the user confirms
        -> confirmed codes go to hs_lookup.search_many() (unchanged)

Design notes
- Digital files (Excel, Word, text PDFs) are read directly: fast + accurate.
- Images and scanned PDFs need OCR (Tesseract). OCR is OPTIONAL: if it isn't
  installed the rest of the app still works and the user gets a clear message.
- Any pattern-matching approach can pick up false positives (invoice numbers,
  dates, phone numbers). That's why every candidate gets a confidence level
  and the user reviews the list before anything is searched.
"""

import csv
import os
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass, field
from typing import Dict, List, Optional


class ExtractionError(Exception):
    """A problem the user can understand and act on (shown in a dialog)."""


# ---------------------------------------------------------------------------
# Public data types
# ---------------------------------------------------------------------------

@dataclass
class Candidate:
    code: str           # digits only, e.g. "85176990"
    confidence: str     # "High" | "Medium" | "Low"
    context: str        # the line / row it was found in (helps the user verify)
    source: str         # "Page 2", "Sheet 'Items'", "Image (OCR)", ...


@dataclass
class ExtractionResult:
    source_name: str
    candidates: List[Candidate] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)


@dataclass
class TextBlock:
    text: str
    source: str
    force_high: bool = False        # True for cells under an "HS Code" column
    context: Optional[str] = None   # what to show the user instead of the raw text


SUPPORTED_EXTENSIONS = {
    ".pdf", ".xlsx", ".xlsm", ".docx", ".csv", ".txt",
    ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp",
}

FILETYPES = [
    ("Supported files",
     "*.pdf *.xlsx *.xlsm *.docx *.csv *.txt *.png *.jpg *.jpeg *.tif *.tiff *.bmp"),
    ("PDF", "*.pdf"),
    ("Excel", "*.xlsx *.xlsm"),
    ("Word", "*.docx"),
    ("Images (needs OCR)", "*.png *.jpg *.jpeg *.tif *.tiff *.bmp"),
    ("All files", "*.*"),
]

OCR_HELP = (
    "Reading images and scanned PDFs needs OCR, which isn't set up on this "
    "computer.\n\n"
    "To enable it:\n"
    "  1. Install Tesseract OCR (Windows installer: "
    "github.com/UB-Mannheim/tesseract/wiki)\n"
    "  2. Run:  pip install pytesseract pillow   (only needed when running "
    "from source)\n\n"
    "Digital PDFs, Excel, Word, CSV and TXT files work without OCR."
)

NUMPY_BROKEN_HELP = (
    "OCR (reading scanned images) needs the 'numpy' package, and the one "
    "installed here is present but broken or incompatible with other packages "
    "already installed - this is a common issue in a mixed conda/pip "
    "environment such as Anaconda's base environment, not a bug in a specific "
    "file being scanned.\n\n"
    "Error detail: {detail}\n\n"
    "Quick fix - reinstall numpy cleanly:\n"
    "  pip install --force-reinstall --no-cache-dir numpy\n\n"
    "If that doesn't resolve it, the most reliable fix is running this app "
    "from its own clean environment instead of the shared Anaconda base one:\n"
    "  conda create -n hslookup python=3.11 -y\n"
    "  conda activate hslookup\n"
    "  pip install -r requirements.txt\n\n"
    "Digital PDFs, Excel, Word, CSV and TXT files don't need numpy or OCR and "
    "work regardless of this."
)


# ---------------------------------------------------------------------------
# Pattern matching
# ---------------------------------------------------------------------------

# Words that signal "the number nearby is an HS code".
KEYWORD = re.compile(
    r"(?<![A-Za-z])(?:hs|hsn|hscode|h\.s\.?|tariff|tarif|harmoni[sz]ed|"
    r"commodity\s+code|customs\s+code)(?![A-Za-z])",
    re.IGNORECASE,
)

# A label sitting DIRECTLY before a number ("Invoice No:", "Tel", "Ref") means
# the number is probably not an HS code. Checking only the text immediately in
# front of the number matters: product names like "Phone parts" or "Videophone"
# are legitimate customs items and must not be penalised.
NOISE_BEFORE = re.compile(
    r"(?<![A-Za-z])(?:inv(?:oice)?|tel(?:ephone)?|phone|mobile|fax|iban|swift|"
    r"account|acc|a/c|po|p\.o\.|order|ref(?:erence)?|date|vat|tin|zip|postal|cr)"
    r"\s*(?:no\.?|num(?:ber)?|#)?\s*[:#.\-]*\s*[+(]?\s*$",
    re.IGNORECASE,
)

# A candidate is either dotted/spaced/dashed groups (8517.69.90, 8517 69 90)
# or 6-12 plain digits. It must not be glued to letters/other digits, and must
# not be the integer part of a decimal like 1234567.89.
#
# The separator must be the SAME throughout (\2 back-reference). Otherwise
# "8517.69.90 10" (a code followed by a quantity of 10) would be glued into a
# fake 10-digit code and the real one would be lost.
CODE_RE = re.compile(
    r"(?<![\w.])"
    r"(\d{4}([.\s-])\d{2}(?:\2\d{2}){0,3}|\d{6,12})"
    r"(?!\w|\.\d)"
)

# "1,234,567" -> "1234567" (thousands separators) without touching lists like
# "85176990,85171100".
_THOUSANDS = re.compile(r"(?<=\d),(?=\d{3}(?!\d))")

# Digit strings that are really dates: YYYYMMDD or DDMMYYYY.
_DATE_YMD = re.compile(r"^(19|20)\d{2}(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])$")
_DATE_DMY = re.compile(r"^(0[1-9]|[12]\d|3[01])(0[1-9]|1[0-2])(19|20)\d{2}$")

_RANK = {"High": 3, "Medium": 2, "Low": 1}


def _clean_context(s: str, limit: int = 110) -> str:
    s = re.sub(r"\s+", " ", s).strip()
    return s if len(s) <= limit else s[: limit - 1] + "\u2026"


def find_hs_codes(blocks: List[TextBlock]) -> List[Candidate]:
    """Scan text blocks and return de-duplicated, ranked HS-code candidates."""
    found: Dict[str, Candidate] = {}
    column_sources = set()     # sources where an "HS Code" column was located

    for block in blocks:
        text = _THOUSANDS.sub("", block.text)
        prev_line = ""
        for line in text.splitlines():
            if not line.strip():
                continue
            kw_same = bool(KEYWORD.search(line))          # "HS Code: 85176990"
            kw_prev = bool(KEYWORD.search(prev_line))     # header row just above

            for m in CODE_RE.finditer(line):
                raw = m.group(1)
                digits = re.sub(r"\D", "", raw)
                n = len(digits)
                spaced = bool(re.search(r"\s", raw))

                if not 1 <= int(digits[:2]) <= 97:      # not a real HS chapter
                    continue

                if block.force_high:
                    confidence = "High"
                else:
                    # Ambiguous shapes are only trusted when the keyword is on
                    # the SAME line (a header line above is too weak a signal).
                    if not kw_same and (_DATE_YMD.match(digits) or _DATE_DMY.match(digits)):
                        continue                          # it's a date
                    if not kw_same and n not in (8, 10, 12):
                        continue                          # 6/7/9/11 digits: too ambiguous

                    # If the number starts the line, its label ("Fax", "Ref"...) may
                    # have been wrapped onto the end of the previous line.
                    before = line[: m.start()]
                    if not before.strip():
                        before = prev_line
                    if NOISE_BEFORE.search(before) and not kw_same:
                        confidence = "Low"                # "Invoice No: 20240512"
                    elif kw_same or (kw_prev and not spaced):
                        confidence = "High"
                    elif spaced:
                        confidence = "Low"                # "1000 50 25" could be a quantity
                    else:
                        confidence = "Medium"

                cand = Candidate(
                    code=digits,
                    confidence=confidence,
                    context=block.context or _clean_context(line),
                    source=block.source,
                )
                if block.force_high:
                    column_sources.add(block.source)
                existing = found.get(digits)
                if existing is None:
                    found[digits] = cand
                elif _RANK[confidence] > _RANK[existing.confidence]:
                    found[digits] = cand                  # keep the better evidence
            prev_line = line

    # If a source has a real "HS Code" column, other numbers that merely LOOK like
    # codes there (product codes, item numbers...) are almost certainly not HS
    # codes: demote them so they are unticked by default on the review screen.
    for cand in found.values():
        if cand.source in column_sources and cand.confidence == "Medium":
            cand.confidence = "Low"

    # High first, then Medium, then Low; original order preserved within a level.
    return sorted(found.values(), key=lambda c: -_RANK[c.confidence])


# ---------------------------------------------------------------------------
# Table handling (shared by Excel, Word, CSV and PDF tables)
# ---------------------------------------------------------------------------

def _cell_text(value) -> str:
    """Excel stores 85176990 as a number (sometimes 85176990.0) - normalise."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _blocks_from_rows(rows, source: str) -> List[TextBlock]:
    """
    Turn table rows into text blocks. If a column header looks like
    'HS Code' / 'Tariff' / 'HSN', every value below it is treated as a
    high-confidence HS code - by far the most reliable signal there is.
    """
    blocks: List[TextBlock] = []
    hs_columns = set()

    for row in rows:
        cells = [_cell_text(c) for c in row]
        row_text = _clean_context(" | ".join(c for c in cells if c))

        for idx, cell in enumerate(cells):
            if cell and len(cell) <= 40 and KEYWORD.search(cell) and not CODE_RE.search(cell):
                hs_columns.add(idx)                       # found a header cell

        for idx, cell in enumerate(cells):
            if not cell:
                continue
            if idx in hs_columns:
                # Excel drops leading zeros on numeric cells (0101.21 -> 10121).
                # HS codes always have an even digit count, so restore it.
                if cell.isdigit() and len(cell) % 2 == 1:
                    cell = "0" + cell
                blocks.append(TextBlock(cell, source, force_high=True, context=row_text))
            else:
                blocks.append(TextBlock(cell, source, context=row_text))
    return blocks


# ---------------------------------------------------------------------------
# OCR (optional)
# ---------------------------------------------------------------------------

def _otsu_threshold(gray: "Image.Image") -> int:
    """Finds the best black/white cutoff for a grayscale image (Otsu's method).
    Pure PIL/Python - no numpy/opencv needed, so OCR stays an optional,
    lightweight extra rather than a heavy dependency."""
    hist = gray.histogram()
    total = sum(hist)
    sum_total = sum(i * h for i, h in enumerate(hist))
    sum_b = 0.0
    w_b = 0
    best_var, best_t = -1.0, 127
    for t in range(256):
        w_b += hist[t]
        if w_b == 0:
            continue
        w_f = total - w_b
        if w_f == 0:
            break
        sum_b += t * hist[t]
        m_b = sum_b / w_b
        m_f = (sum_total - sum_b) / w_f
        var_between = w_b * w_f * (m_b - m_f) ** 2
        if var_between > best_var:
            best_var, best_t = var_between, t
    return best_t


def _deskew(gray: "Image.Image", pytesseract) -> "Image.Image":
    """Straightens rotated scans/photos using Tesseract's own orientation
    detector. A tilted page is a major source of misread digits (a skewed
    '0' can look like a '9' to the OCR engine) - this is not optional
    polish, it measurably changes which digits come out correct."""
    try:
        osd = pytesseract.image_to_osd(gray)
    except Exception:
        return gray  # OSD needs enough text to work with; skip quietly if it fails
    m = re.search(r"Rotate:\s*(\d+)", osd)
    if not m:
        return gray
    angle = int(m.group(1))
    if angle:
        gray = gray.rotate(-angle, expand=True, fillcolor=255)
    return gray


def _fine_deskew(bw: "Image.Image") -> "Image.Image":
    """Corrects small tilt angles (a few degrees) typical of a slightly
    crooked phone photo - Tesseract's own OSD (used in _deskew above) only
    detects 90/180/270-degree rotations, not this. Finds the angle where
    horizontal text lines line up most cleanly (their row-brightness profile
    becomes most 'peaky': solid text rows vs. solid gaps between lines).

    This is a best-effort enhancement, not a requirement: if numpy isn't
    installed, or is installed but broken (a mixed pip/conda environment can
    produce a binary-incompatibility error the moment numpy is imported, not
    just a missing-package error), this quietly skips fine correction rather
    than failing the whole scan. The coarse OSD correction above still runs
    either way, so rotated documents are never left completely unhandled."""
    try:
        import numpy as np

        thumb = bw.copy()
        thumb.thumbnail((700, 700))  # search on a small copy for speed
        best_angle, best_score = 0.0, -1.0
        for tenth_degree in range(-150, 151, 5):    # -15.0 to +15.0 in 0.5-degree steps
            angle = tenth_degree / 10
            rotated = thumb.rotate(angle, expand=True, fillcolor=255)
            rows = (np.array(rotated) < 128).sum(axis=1)  # black-pixel count per row
            score = float(rows.var())
            if score > best_score:
                best_score, best_angle = score, angle
        if abs(best_angle) >= 0.5:
            return bw.rotate(best_angle, expand=True, fillcolor=255)
        return bw
    except Exception:
        # Anything here (missing numpy, a broken/incompatible numpy install,
        # an unexpected array shape) - fall back to the un-fine-tuned image
        # rather than losing the whole document scan over one enhancement step.
        return bw


def _bundled_tesseract():
    """Finds a copy of Tesseract shipped WITH the app (a 'tesseract' folder that
    build.bat packs into the .exe), so the client installs nothing.
    Returns (path_to_exe, path_to_tessdata) or None."""
    roots = []
    if getattr(sys, "_MEIPASS", None):                       # PyInstaller build
        roots.append(sys._MEIPASS)
    roots.append(os.path.dirname(os.path.abspath(__file__)))  # running from source
    roots.append(os.path.dirname(os.path.abspath(sys.executable)))  # next to the .exe
    for root in roots:
        folder = os.path.join(root, "tesseract")
        for name in ("tesseract.exe", "tesseract"):
            exe = os.path.join(folder, name)
            if os.path.isfile(exe):
                return exe, os.path.join(folder, "tessdata")
    return None


_OCR_CHECKED: Dict[str, str] = {}      # tesseract command -> "" if it starts fine, else the reason


def _tesseract_self_test(cmd: str) -> str:
    """Actually starts Tesseract once ('--version') and returns "" if it works,
    otherwise the real reason it doesn't. pytesseract reports EVERY failure to
    launch (file missing, missing Windows runtime DLL, blocked by antivirus...)
    as the same 'not found' error, which hides what is really wrong."""
    if cmd in _OCR_CHECKED:
        return _OCR_CHECKED[cmd]
    reason = ""
    try:
        flags = 0x08000000 if os.name == "nt" else 0          # no console window flash
        proc = subprocess.run([cmd, "--version"], capture_output=True, timeout=30,
                              creationflags=flags)
        if proc.returncode != 0:
            err = (proc.stderr or proc.stdout or b"").decode("utf-8", "replace").strip()
            reason = f"exit code {proc.returncode}. {err[:200]}"
    except Exception as e:                                     # OSError, timeout, ...
        reason = f"{type(e).__name__}: {e}"
    _OCR_CHECKED[cmd] = reason
    return reason


def _ocr_problem_message(bundled, cmd: str, reason: str) -> str:
    if bundled:
        return (
            "The OCR engine built into this app could not start on this computer.\n\n"
            f"Technical detail: {reason}\n\n"
            "Most common fix: install the free 'Microsoft Visual C++ Redistributable "
            "2015-2022 (x64)' from Microsoft (aka.ms/vs/17/release/vc_redist.x64.exe) "
            "and restart the app.\n"
            "If it still fails, antivirus software may be blocking the OCR engine - "
            "allow this app in the antivirus."
        )
    extra = ("\n\n(This copy of the app was built WITHOUT the built-in OCR engine - "
             "the 'tesseract' folder was missing when it was built.)"
             if getattr(sys, "frozen", False) else "")
    return OCR_HELP + extra + f"\n\nTechnical detail: {reason}"


def _to_rgb_on_white(img):
    """Any image mode -> plain RGB (transparent areas become white, not black)."""
    from PIL import Image
    if img.mode in ("RGBA", "LA", "P"):
        rgba = img.convert("RGBA")
        bg = Image.new("RGB", rgba.size, (255, 255, 255))
        bg.paste(rgba, mask=rgba.split()[3])
        return bg
    return img.convert("RGB")


def _remove_colored_ink(img):
    """Returns a copy with strongly coloured pixels (pen/marker scribbles,
    highlighter, stamps) painted white, or None if there is nothing to remove.

    Why: binarisation turns a blue pen scribble into a solid black blob, which
    wrecks Tesseract's line detection for the whole table row it crosses (the
    row's digits come out as garbage or vanish). Printed text is black/grey
    (low saturation), so removing high-saturation pixels removes the ink and
    keeps the text."""
    try:
        rgb = _to_rgb_on_white(img)
        sat = rgb.convert("HSV").split()[1]
        mask = sat.point(lambda v: 255 if v > 110 else 0)
        share = mask.histogram()[255] / float(rgb.width * rgb.height)
        if share < 0.0005:              # < 0.05% coloured pixels: nothing to clean
            return None
        cleaned = rgb.copy()
        cleaned.paste((255, 255, 255), mask=mask)
        return cleaned
    except Exception:
        return None


# --- Column-aware reading ---------------------------------------------------
# Plain OCR text loses the table layout. Here we keep word positions, find the
# "HS Code" header, and read only the numbers directly underneath it. That is
# what tells the HS Code column apart from Product Code / Item No. columns.

_HS_HEADER_WORD = re.compile(r"^(hs|hsn|hscode|tariff|tarif)$")
_STOP_WORD = re.compile(r"^(grand|total|subtotal|currency)$")
_NUMERIC_TOKEN = re.compile(r"^[|\[\]()]*\d[\d.\-]*[|\[\]()]*$")


def _letters(s: str) -> str:
    return re.sub(r"[^a-z]", "", s.lower())


def _group_rows(words):
    heights = sorted(w["h"] for w in words)
    tol = max(4.0, 0.6 * heights[len(heights) // 2])
    rows = []
    for w in sorted(words, key=lambda w: w["cy"]):
        if rows and abs(w["cy"] - rows[-1]["cy"]) <= tol:
            rows[-1]["words"].append(w)
            ws = rows[-1]["words"]
            rows[-1]["cy"] = sum(x["cy"] for x in ws) / len(ws)
        else:
            rows.append({"cy": w["cy"], "words": [w]})
    for r in rows:
        r["words"].sort(key=lambda w: w["l"])
    return rows


def _in_band(w, lo, hi) -> bool:
    overlap = min(w["r"], hi) - max(w["l"], lo)
    return overlap >= 0.5 * max(1, w["r"] - w["l"])


def _codes_under_headers(rows):
    out = []
    for ri, row in enumerate(rows):
        ws = row["words"]
        # A real table header row has no long numbers in it ("HS Code: 8517..."
        # on one line is handled by the normal keyword logic instead).
        if any(len(re.sub(r"\D", "", w["t"])) >= 6 for w in ws):
            continue
        for wi, w in enumerate(ws):
            if not _HS_HEADER_WORD.match(_letters(w["t"])):
                continue
            left, right = w["l"], w["r"]
            if (_letters(w["t"]) == "hs" and wi + 1 < len(ws)
                    and _letters(ws[wi + 1]["t"]) in ("code", "no")):
                right = ws[wi + 1]["r"]              # "HS" "Code" as two words
            pad = 0.35 * (right - left)
            lo, hi = left - pad, right + pad

            for below in rows[ri + 1:]:
                if any(_STOP_WORD.match(_letters(x["t"])) for x in below["words"]):
                    break                            # reached the totals
                toks = [x["t"] for x in below["words"]
                        if _in_band(x, lo, hi) and _NUMERIC_TOKEN.match(x["t"])]
                if not toks:
                    continue
                joined = "".join(re.sub(r"[^\d.\-]", "", t) for t in toks)
                # Skip decimals like 527.000 / weights; keep 8517.69.90 style.
                if joined.count(".") == 1:
                    ip, fp = joined.split(".")
                    if not (len(ip) == 4 and len(fp) == 2):
                        continue
                digits = re.sub(r"\D", "", joined)
                if 6 <= len(digits) <= 12:
                    ctx = _clean_context(" | ".join(x["t"] for x in below["words"]))
                    out.append((digits, ctx))
    return out


def _column_codes(bw, pytesseract):
    for psm in (6, 4):
        try:
            d = pytesseract.image_to_data(
                bw, config=f"--psm {psm}", output_type=pytesseract.Output.DICT
            )
        except Exception:
            continue
        words = []
        for i, t in enumerate(d["text"]):
            t = (t or "").strip()
            try:
                conf = float(d["conf"][i])
            except (TypeError, ValueError):
                conf = -1
            if not t or conf < 0:
                continue
            l, top, w, h = d["left"][i], d["top"][i], d["width"][i], d["height"][i]
            words.append({"t": t, "l": l, "r": l + w, "cy": top + h / 2.0, "h": h})
        if not words:
            continue
        codes = _codes_under_headers(_group_rows(words))
        if codes:
            return codes
    return []


# --- Main OCR entry points --------------------------------------------------

def _ocr_variant(base, pytesseract, ImageOps, ImageFilter):
    """OCR one version of the image. Returns (text, [(code, row_context), ...])."""
    img = base.convert("L")  # grayscale

    # Cap runaway resolution (very large phone photos) before the expensive
    # steps below, then upscale small/low-res images - both directions hurt
    # OCR accuracy on this kind of small, dense text.
    if img.width > 3500:
        ratio = 3500 / img.width
        img = img.resize((3500, int(img.height * ratio)))
    if img.width < 2000:
        factor = max(2, round(2000 / max(img.width, 1)))
        img = img.resize((img.width * factor, img.height * factor))

    img = ImageOps.autocontrast(img, cutoff=1)          # fixes washed-out/low-contrast scans
    img = img.filter(ImageFilter.UnsharpMask(radius=2, percent=150))
    # Binarize BEFORE any rotation of our own. Otsu's split works out badly if
    # the image already contains white corner-padding from an earlier rotation
    # (it starts separating "our padding" from "the real page" instead of
    # "background" from "text") - so threshold first, then rotate the already-
    # binary image, where new padding is safely just more background (255).
    threshold = _otsu_threshold(img)
    bw = img.point(lambda p: 255 if p > threshold else 0)
    bw = _deskew(bw, pytesseract)     # coarse: fixes 90/180/270-degree rotation
    bw = _fine_deskew(bw)             # fine: fixes a few degrees of tilt

    # Run two passes with different page-segmentation assumptions and merge
    # the text: PSM 6 is best for a solid block of text, PSM 11 is best for
    # scattered/sparse lines (labels dotted around a photo). Combining both
    # catches more real codes than betting on a single mode.
    outputs = []
    for psm in (6, 11):
        try:
            outputs.append(pytesseract.image_to_string(bw, config=f"--psm {psm}"))
        except pytesseract.TesseractNotFoundError:
            raise ExtractionError(OCR_HELP)
        except pytesseract.TesseractError:
            continue
    return "\n".join(outputs), _column_codes(bw, pytesseract)


def _ocr_full(pil_image):
    """OCR an image. Returns (text, column_codes) where column_codes are numbers
    found directly under an 'HS Code' table header: [(code, row_text), ...]."""
    try:
        import pytesseract
        from PIL import ImageOps, ImageFilter
    except ImportError as e:
        raise ExtractionError(OCR_HELP + f"\n\nTechnical detail: {e}")
    except Exception as e:
        # pytesseract imports numpy internally and only guards against numpy
        # being MISSING (ModuleNotFoundError) - a numpy that's installed but
        # binary-incompatible with other packages raises a different error
        # (ValueError) that slips past that guard. Without this, the person
        # would see a cryptic "numpy.dtype size changed..." message instead
        # of an actionable one.
        raise ExtractionError(NUMPY_BROKEN_HELP.format(detail=e))

    # 1) Prefer the Tesseract bundled inside the app (nothing to install).
    bundled = _bundled_tesseract()
    if bundled:
        pytesseract.pytesseract.tesseract_cmd = bundled[0]
        os.environ["TESSDATA_PREFIX"] = bundled[1]
    # 2) Otherwise a system install. Windows: the installer doesn't always add
    #    itself to PATH.
    elif not shutil.which("tesseract"):
        for p in (
            r"C:\Program Files\Tesseract-OCR\tesseract.exe",
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        ):
            if os.path.exists(p):
                pytesseract.pytesseract.tesseract_cmd = p
                break

    # Make sure Tesseract can really start, and say exactly why if it can't.
    cmd = pytesseract.pytesseract.tesseract_cmd
    reason = _tesseract_self_test(cmd)
    if reason:
        raise ExtractionError(_ocr_problem_message(bundled, cmd, reason))

    # Phone photos carry an EXIF "orientation" tag saying how to rotate the
    # image for display, but Pillow (and therefore OCR) reads the RAW pixels
    # unless told otherwise. Without this, a portrait phone photo read into
    # a sideways image and OCR finds nothing at all.
    base = _to_rgb_on_white(ImageOps.exif_transpose(pil_image))

    # If the page has pen/marker scribbles, read a cleaned copy AND the original
    # (in case the coloured pixels are real printed text) and merge the results.
    variants = []
    cleaned = _remove_colored_ink(base)
    if cleaned is not None:
        variants.append(cleaned)
    variants.append(base)

    texts, columns = [], []
    for v in variants:
        text, cols = _ocr_variant(v, pytesseract, ImageOps, ImageFilter)
        texts.append(text)
        for item in cols:
            if item[0] not in [c[0] for c in columns]:
                columns.append(item)
    return "\n".join(texts), columns


def _ocr_image(pil_image) -> str:
    """Text only (kept for callers that don't need the column information)."""
    return _ocr_full(pil_image)[0]


# ---------------------------------------------------------------------------
# Readers, one per file type
# ---------------------------------------------------------------------------

def _read_text_file(path: str) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            with open(path, encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Couldn't read that text file (unknown encoding).")


def _from_txt(path):
    return [TextBlock(_read_text_file(path), "Text file")], []


def _from_csv(path):
    text = _read_text_file(path)
    rows = list(csv.reader(text.splitlines()))
    return _blocks_from_rows(rows, "CSV"), []


def _from_excel(path):
    try:
        import openpyxl
    except ImportError:
        raise ExtractionError("Excel support needs the 'openpyxl' package (pip install openpyxl).")
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise ExtractionError(f"Couldn't open that Excel file: {e}")
    blocks = []
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            blocks.extend(_blocks_from_rows(rows, f"Sheet '{ws.title}'"))
    finally:
        wb.close()
    return blocks, []


def _from_docx(path):
    try:
        import docx
    except ImportError:
        raise ExtractionError("Word support needs the 'python-docx' package (pip install python-docx).")
    try:
        document = docx.Document(path)
    except Exception as e:
        raise ExtractionError(f"Couldn't open that Word file: {e}")
    blocks = [TextBlock("\n".join(p.text for p in document.paragraphs), "Document text")]
    for t_idx, table in enumerate(document.tables, 1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        blocks.extend(_blocks_from_rows(rows, f"Table {t_idx}"))
    return blocks, []


def _from_pdf(path):
    try:
        import pdfplumber
    except ImportError:
        raise ExtractionError("PDF support needs the 'pdfplumber' package (pip install pdfplumber).")

    blocks: List[TextBlock] = []
    warnings: List[str] = []
    scanned_without_ocr = 0
    last_ocr_error = ""

    try:
        with pdfplumber.open(path) as pdf:
            for page_no, page in enumerate(pdf.pages, 1):
                source = f"Page {page_no}"

                try:
                    for table in page.extract_tables():
                        blocks.extend(_blocks_from_rows(table, source))
                except Exception:
                    pass                                   # table detection is a bonus

                text = page.extract_text() or ""
                if text.strip():
                    blocks.append(TextBlock(text, source))
                    continue

                # No text layer -> it's a scanned page; fall back to OCR.
                try:
                    img = page.to_image(resolution=200).original
                    ocr_text, ocr_cols = _ocr_full(img)
                    blocks.append(TextBlock(ocr_text, f"{source} (OCR)"))
                    for code, ctx in ocr_cols:
                        blocks.append(TextBlock(code, f"{source} (OCR)",
                                                force_high=True, context=ctx))
                except ExtractionError as e:
                    scanned_without_ocr += 1
                    last_ocr_error = str(e)
    except ExtractionError:
        raise
    except Exception as e:
        raise ExtractionError(f"Couldn't read that PDF: {e}")

    if scanned_without_ocr:
        if not blocks:
            raise ExtractionError(
                "This PDF looks like a scan (no selectable text).\n\n"
                + (last_ocr_error or OCR_HELP)
            )
        warnings.append(
            f"{scanned_without_ocr} scanned page(s) were skipped because OCR isn't working."
        )
    elif any("(OCR)" in b.source for b in blocks):
        warnings.append("Some pages were read with OCR - double-check those codes carefully.")
    return blocks, warnings


def _from_image(path):
    try:
        from PIL import Image
    except ImportError:
        raise ExtractionError(OCR_HELP)
    try:
        img = Image.open(path)
    except Exception as e:
        raise ExtractionError(f"Couldn't open that image: {e}")
    text, cols = _ocr_full(img)
    blocks = [TextBlock(text, "Image (OCR)")]
    for code, ctx in cols:
        blocks.append(TextBlock(code, "Image (OCR)", force_high=True, context=ctx))
    return blocks, ["Read with OCR - double-check each code before searching."]


_READERS = {
    ".txt": _from_txt,
    ".csv": _from_csv,
    ".xlsx": _from_excel,
    ".xlsm": _from_excel,
    ".docx": _from_docx,
    ".pdf": _from_pdf,
    ".png": _from_image, ".jpg": _from_image, ".jpeg": _from_image,
    ".tif": _from_image, ".tiff": _from_image, ".bmp": _from_image,
}


# ---------------------------------------------------------------------------
# Entry point used by the GUI
# ---------------------------------------------------------------------------

def extract_candidates(path: str) -> ExtractionResult:
    ext = os.path.splitext(path)[1].lower()
    name = os.path.basename(path)

    if ext in (".xls", ".doc"):
        raise ExtractionError(
            f"Old-format '{ext}' files aren't supported. Open it and use "
            f"'Save As' to convert it to {'.xlsx' if ext == '.xls' else '.docx'}."
        )
    reader = _READERS.get(ext)
    if reader is None:
        raise ExtractionError(
            "That file type isn't supported.\nSupported: PDF, Excel (.xlsx), Word (.docx), "
            "CSV, TXT, and images (PNG/JPG/TIFF/BMP)."
        )

    blocks, warnings = reader(path)
    return ExtractionResult(
        source_name=name,
        candidates=find_hs_codes(blocks),
        warnings=warnings,
    )